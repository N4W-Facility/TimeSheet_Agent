# ============================================================
# N4W FACILITY — Excel formato PowerApp + entrega en OneDrive
# ============================================================
import logging
import os
import re
import shutil
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import N4W_ONEDRIVE_FOLDER
from core.utils import get_date_columns, is_special_code

try:
    import winreg  # solo Windows
except ImportError:
    winreg = None

log = logging.getLogger(__name__)

TABLE_NAME = 'new_n4wtimeentriessubmissionses'
HEADERS = [
    'new_title', 'new_employeeemail', 'new_employeename', 'new_projectcode',
    'new_monhours', 'new_tuehours', 'new_wedhours', 'new_thurshours',
    'new_frihours', 'new_sathours', 'new_sunhours', 'new_totalhours',
    'crd63_timesheetinitiated', 'new_timesheetstatus', 'crd63_weekstartdate', 'new_comments',
]
COLUMN_WIDTHS = [10.08, 19.18, 19.27, 16.63, 15.00, 14.09, 14.73, 15.63,
                 13.09, 14.00, 14.36, 15.18, 23.09, 20.18, 19.91, 15.54]


# ── Excel N4W ────────────────────────────────────────────────

def load_timesheet_codes(task_details_path: str) -> Dict[str, str]:
    """Task_Name → Timesheet Code (hoja Task_Details)."""
    df = pd.read_excel(task_details_path, sheet_name='Task_Details')
    return dict(zip(df['Task_Name'], df['Timesheet Code']))


def build_n4w_rows(timesheet: pd.DataFrame, email: str, name: str,
                   timesheet_codes: Dict[str, str]) -> pd.DataFrame:
    """Una fila por (proyecto, semana lunes–domingo) con horas > 0."""
    df = timesheet.copy()
    df.loc[df['Code'].map(is_special_code), 'Code'] = 'OF0104'

    date_cols = [c for c in get_date_columns(df) if '00:00:00' in c]
    dates = [datetime.strptime(c.replace(' 00:00:00', ''), '%Y-%m-%d') for c in date_cols]

    df = df[~df['Code'].astype(str).str.startswith('TNC')]
    df = df.groupby(['Code'], as_index=False)[date_cols].sum()

    day_keys = ['sun', 'mon', 'tue', 'wed', 'thu', 'fri', 'sat']
    rows = []
    for _, row in df.iterrows():
        code = row['Code']
        project_code = timesheet_codes.get(code, code)
        if timesheet_codes and code not in timesheet_codes:
            log.warning(f"No Timesheet Code for {code}; using original code")

        by_week = {}
        for col, d in zip(date_cols, dates):
            week_start = d - timedelta(days=d.weekday())   # lunes
            by_week.setdefault(week_start, dict.fromkeys(day_keys, 0))
            by_week[week_start][day_keys[(d.weekday() + 1) % 7]] += row[col] if pd.notna(row[col]) else 0

        for week_start, h in by_week.items():
            total = sum(h.values())
            if total <= 0:
                continue
            week_end = week_start + timedelta(days=6)
            rows.append({
                'new_title': f"{week_start.strftime('%d-%B-%Y')} to {week_end.strftime('%d-%B-%Y')}",
                'new_employeeemail': email,
                'new_employeename': name,
                'new_projectcode': project_code,
                'new_monhours': h['mon'], 'new_tuehours': h['tue'], 'new_wedhours': h['wed'],
                'new_thurshours': h['thu'], 'new_frihours': h['fri'], 'new_sathours': h['sat'],
                'new_sunhours': h['sun'],
                'new_totalhours': total,
                'crd63_timesheetinitiated': True,
                'new_timesheetstatus': 'Submitted',
                'crd63_weekstartdate': week_start,
                'new_comments': 'Submitted',
            })
    out = pd.DataFrame(rows, columns=HEADERS)
    return out.sort_values('crd63_weekstartdate')


def write_n4w_excel(rows: pd.DataFrame, output_path: str) -> str:
    """Escribe el Excel con la tabla que espera el flujo de PowerApp."""
    wb = Workbook()
    ws = wb.active
    ws.title = TABLE_NAME

    for col, header in enumerate(HEADERS, 1):
        ws.cell(row=1, column=col, value=header)

    for r, (_, row) in enumerate(rows.iterrows(), 2):
        for c, header in enumerate(HEADERS, 1):
            value = row[header]
            cell = ws.cell(row=r, column=c, value=value)
            if c == 15:                       # fecha inicio semana
                cell.value = to_excel(value)
                cell.number_format = 'M/D/YY'
            elif 5 <= c <= 12:                # horas: vacío si 0
                cell.value = None if value == 0 else float(value)
                cell.number_format = 'General'
            elif c == 13:
                cell.value = True
                cell.number_format = 'General'
            else:
                cell.number_format = 'General'

    if len(rows) > 0:
        ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
        table = Table(displayName=TABLE_NAME, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                              showLastColumn=False, showRowStripes=True,
                                              showColumnStripes=True)
        ws.add_table(table)

    for col, width in enumerate(COLUMN_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(output_path)
    log.info(f"N4W Excel created: {output_path} ({len(rows)} rows)")
    return output_path


def n4w_summary(rows: pd.DataFrame) -> str:
    """Tabla de texto con lo que se envía: una línea por (semana, proyecto) y el total."""
    days = [('Mon', 'new_monhours'), ('Tue', 'new_tuehours'), ('Wed', 'new_wedhours'),
            ('Thu', 'new_thurshours'), ('Fri', 'new_frihours'), ('Sat', 'new_sathours'),
            ('Sun', 'new_sunhours')]
    h = lambda v: f"{float(v):g}" if v else "·"
    width = max([7] + [len(str(c)) for c in rows['new_projectcode']])
    lines = [f"{'Week':<10}  {'Project':<{width}}  " + " ".join(f"{d:>4}" for d, _ in days)
             + f"  {'Total':>5}"]
    for _, r in rows.iterrows():
        lines.append(f"{pd.Timestamp(r['crd63_weekstartdate']):%Y-%m-%d}  "
                     f"{str(r['new_projectcode']):<{width}}  "
                     + " ".join(f"{h(r[c]):>4}" for _, c in days)
                     + f"  {h(r['new_totalhours']):>5}")
    lines.append(f"{'':<10}  {'TOTAL':<{width}}  {'':<{5 * len(days) - 1}}"
                 f"  {h(rows['new_totalhours'].sum()):>5}")
    return "\n".join(lines)


def n4w_filename(email: str, start: datetime, end: datetime) -> str:
    return f"{email}_{start.strftime('%Y-%m-%d')}_to_{end.strftime('%Y-%m-%d')}.xlsx"


# ── OneDrive ─────────────────────────────────────────────────

def _env_onedrive_candidates() -> List[Path]:
    seen, out = set(), []
    for var in ("OneDriveCommercial", "OneDriveConsumer", "OneDrive"):
        p = os.environ.get(var)
        if p:
            pp = Path(p).expanduser().resolve()
            if pp.exists() and pp not in seen:
                out.append(pp)
                seen.add(pp)
    return out


def _reg_get_str(key, value_name: str) -> Optional[str]:
    try:
        val, _ = winreg.QueryValueEx(key, value_name)
        return val if isinstance(val, str) else None
    except OSError:
        return None


def _registry_onedrive_folders() -> List[str]:
    if winreg is None:
        return []
    folders = []
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\OneDrive\Accounts") as accounts:
            i = 0
            while True:
                try:
                    subname = winreg.EnumKey(accounts, i)
                    i += 1
                except OSError:
                    break
                try:
                    with winreg.OpenKey(accounts, subname) as subkey:
                        folder = _reg_get_str(subkey, "UserFolder")
                        if folder:
                            folders.append(folder)
                except OSError:
                    continue
    except OSError:
        pass
    return folders


def get_onedrive_accounts() -> List[Dict[str, str]]:
    """[{'label': 'OneDrive - Empresa', 'root': 'C:\\Users\\...'}] desde registro + variables de entorno."""
    roots = [Path(f).expanduser().resolve() for f in _registry_onedrive_folders()]
    roots += _env_onedrive_candidates()
    seen, out = set(), []
    for r in roots:
        if r.exists() and str(r) not in seen:
            out.append({"label": r.name, "root": str(r)})
            seen.add(str(r))
    return out


def resolve_onedrive_target(target_path: str, account_hint: Optional[str] = None) -> Path:
    """
    Ruta absoluta destino. Nota: la carpeta compartida vive bajo
    '<padre de OneDrive>/The Nature Conservancy' (biblioteca sincronizada de SharePoint).
    """
    parts = [p for p in re.split(r"[\\/]+", target_path.strip().strip("\\/")) if p]
    if not parts:
        raise ValueError("Empty destination path.")

    accounts = get_onedrive_accounts()
    if not accounts:
        raise RuntimeError("No OneDrive folder detected in this Windows profile.")

    by_label = {a["label"]: a for a in accounts}
    if parts[0] in by_label:
        return Path(by_label[parts[0]]["root"]).joinpath(*parts[1:])

    chosen = None
    if len(accounts) == 1:
        chosen = accounts[0]
    elif account_hint:
        hint = account_hint.lower()
        candidates = [a for a in accounts if hint in a["label"].lower()]
        if candidates:
            candidates.sort(key=lambda a: a["label"].lower().find(hint))
            chosen = candidates[0]
    if not chosen:
        enterprise = [a for a in accounts if " - " in a["label"]]
        chosen = enterprise[0] if enterprise else accounts[0]

    root = Path(os.path.split(chosen["root"])[0]) / "The Nature Conservancy"
    return root.joinpath(*parts)


def put_file_in_onedrive(src_path: str, target_path: str, account_hint: Optional[str] = None,
                         overwrite: bool = False) -> str:
    """Copia un archivo a la carpeta compartida de OneDrive/SharePoint."""
    src = Path(src_path).expanduser().resolve()
    if not src.is_file():
        raise FileNotFoundError(f"Source file does not exist: {src}")

    dst = resolve_onedrive_target(target_path, account_hint=account_hint)
    dst.parent.mkdir(parents=True, exist_ok=True)

    if dst.exists():
        if not overwrite:
            raise FileExistsError(f"Destination already exists: {dst}")
        if not dst.is_file():
            raise IsADirectoryError(f"Destination is a folder: {dst}")
        try:
            os.remove(dst)
        except PermissionError:   # bloqueado por sincronización → renombrar
            dst.replace(dst.with_suffix(dst.suffix + f".bak.{uuid.uuid4().hex[:8]}"))

    # Comportamiento heredado: se quita 'OneDrive - ' / 'OneDrive' de la ruta final
    final = str(dst).replace('OneDrive - ', '').replace('OneDrive', '')
    shutil.copy2(str(src), final)
    return final


# ── Duplicados ───────────────────────────────────────────────

def parse_filename_dates(filename: str) -> tuple:
    m = re.match(r'.*_(\d{4}-\d{2}-\d{2})_to_(\d{4}-\d{2}-\d{2})\.xlsx$', filename)
    if not m:
        return None, None
    return tuple(datetime.strptime(s, '%Y-%m-%d').date() for s in m.groups())


def find_existing_submissions(email: str) -> List[tuple]:
    """[(filename, start, end)] ya entregados en la carpeta N4W de OneDrive."""
    folder = resolve_onedrive_target(N4W_ONEDRIVE_FOLDER)
    if not folder.exists():
        log.warning(f"N4W folder not found in OneDrive: {folder}")
        return []
    out = []
    for f in folder.glob(f"{email}_*.xlsx"):
        s, e = parse_filename_dates(f.name)
        if s and e:
            out.append((f.name, s, e))
    return out


def find_overlaps(new_start, new_end, existing: List[tuple]) -> List[dict]:
    """Entregas existentes que se solapan con [new_start, new_end]."""
    ns = new_start.date() if hasattr(new_start, 'date') else new_start
    ne = new_end.date() if hasattr(new_end, 'date') else new_end
    return [{'filename': f, 'start': s, 'end': e}
            for f, s, e in existing if not (ne < s or ns > e)]
